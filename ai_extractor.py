"""
Module d'extraction de données à partir de documents avec IA
"""
import os
import json
import re
from typing import Dict, Any, List
import openai
from PIL import Image
import pytesseract
from pdf2image import convert_from_path
import PyPDF2
from docx import Document

class DocumentExtractor:
    def __init__(self, openai_api_key: str = None):
        """Initialise l'extracteur avec la clé API OpenAI"""
        if openai_api_key:
            openai.api_key = openai_api_key
        
    def extract_text_from_pdf(self, file_path: str) -> str:
        """Extrait le texte d'un fichier PDF"""
        text = ""
        try:
            # Méthode 1: Extraction directe du texte
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
            
            # Si peu de texte extrait, utiliser OCR
            if len(text.strip()) < 100:
                images = convert_from_path(file_path)
                for image in images:
                    text += pytesseract.image_to_string(image, lang='fra') + "\n"
                    
        except Exception as e:
            print(f"Erreur lors de l'extraction PDF: {e}")
            # Fallback sur OCR
            try:
                images = convert_from_path(file_path)
                for image in images:
                    text += pytesseract.image_to_string(image, lang='fra') + "\n"
            except Exception as e2:
                print(f"Erreur OCR: {e2}")
                
        return text
    
    def extract_text_from_docx(self, file_path: str) -> str:
        """Extrait le texte d'un fichier Word"""
        try:
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Extraire aussi le contenu des tableaux
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + "\t"
                    text += "\n"
                    
            return text
        except Exception as e:
            print(f"Erreur lors de l'extraction DOCX: {e}")
            return ""
    
    def extract_text_from_file(self, file_path: str) -> str:
        """Extrait le texte selon le type de fichier"""
        if file_path.lower().endswith('.pdf'):
            return self.extract_text_from_pdf(file_path)
        elif file_path.lower().endswith(('.docx', '.doc')):
            return self.extract_text_from_docx(file_path)
        else:
            return ""
    
    def extract_data_with_ai(self, text: str, document_type: str) -> Dict[str, Any]:
        """Utilise l'IA pour extraire des données structurées"""
        
        # Prompts spécialisés selon le type de document
        prompts = {
            'kbis': """
            Analyse ce texte d'extrait Kbis et extrait les informations suivantes au format JSON:
            {
                "denomination_sociale": "",
                "capital_social": "",
                "adresse_siege": "",
                "numero_rcs": "",
                "ville_greffe": "",
                "date_immatriculation": "",
                "dirigeants": []
            }
            """,
            
            'statuts': """
            Analyse ce texte de statuts de SAS et extrait les informations suivantes au format JSON:
            {
                "denomination_sociale": "",
                "objet_social": "",
                "duree_societe": "",
                "capital_social": "",
                "siege_social": "",
                "president": "",
                "exercice_social": ""
            }
            """,
            
            'comptes': """
            Analyse ces comptes annuels et extrait les données financières au format JSON:
            {
                "exercice": "",
                "chiffre_affaires": "",
                "resultat_net": "",
                "total_bilan": "",
                "capitaux_propres": "",
                "dettes": "",
                "immobilisations": ""
            }
            """
        }
        
        if document_type not in prompts:
            return {}
        
        try:
            # Appel à l'API OpenAI (à adapter selon la version utilisée)
            response = openai.ChatCompletion.create(
                model="gpt-3.5-turbo",
                messages=[
                    {"role": "system", "content": "Tu es un expert en analyse de documents juridiques français."},
                    {"role": "user", "content": f"{prompts[document_type]}\n\nTexte à analyser:\n{text[:4000]}"}
                ],
                temperature=0.1
            )
            
            # Extraire et parser la réponse JSON
            content = response.choices[0].message.content
            # Chercher le JSON dans la réponse
            json_match = re.search(r'\{.*\}', content, re.DOTALL)
            if json_match:
                return json.loads(json_match.group())
            else:
                return {}
                
        except Exception as e:
            print(f"Erreur IA pour {document_type}: {e}")
            # Fallback: extraction par regex basique
            return self._fallback_extraction(text, document_type)
    
    def _fallback_extraction(self, text: str, document_type: str) -> Dict[str, Any]:
        """Extraction basique par regex en cas d'échec de l'IA"""
        data = {}
        
        if document_type == 'kbis':
            # Recherche de patterns typiques dans un Kbis
            siren_match = re.search(r'SIREN\s*:?\s*(\d{9})', text, re.IGNORECASE)
            if siren_match:
                data['numero_rcs'] = siren_match.group(1)
                
            capital_match = re.search(r'capital.*?(\d+(?:[\s,\.]\d+)*)\s*euros?', text, re.IGNORECASE)
            if capital_match:
                data['capital_social'] = capital_match.group(1)
        
        elif document_type == 'statuts':
            # Recherche dans les statuts
            denomination_match = re.search(r'dénomination\s*:?\s*(.+?)(?:\n|,)', text, re.IGNORECASE)
            if denomination_match:
                data['denomination_sociale'] = denomination_match.group(1).strip()
        
        elif document_type == 'comptes':
            # Recherche de montants financiers
            ca_match = re.search(r'chiffre.*?affaires.*?(\d+(?:[\s,\.]\d+)*)', text, re.IGNORECASE)
            if ca_match:
                data['chiffre_affaires'] = ca_match.group(1)
        
        return data
    
    def process_documents(self, files_data: Dict[str, str]) -> Dict[str, Any]:
        """Traite tous les documents et retourne les données extraites"""
        extracted_data = {}
        
        # Mapping des types de fichiers
        file_types = {
            'kbis': 'kbis',
            'statuts': 'statuts',
            'comptes_n1': 'comptes',
            'comptes_n2': 'comptes',
            'comptes_n3': 'comptes'
        }
        
        for file_key, file_path in files_data.items():
            if file_key in file_types:
                print(f"Traitement de {file_key}: {file_path}")
                
                # Extraction du texte
                text = self.extract_text_from_file(file_path)
                
                if text.strip():
                    # Extraction des données structurées
                    doc_type = file_types[file_key]
                    data = self.extract_data_with_ai(text, doc_type)
                    extracted_data[file_key] = data
                else:
                    print(f"Aucun texte extrait de {file_path}")
                    extracted_data[file_key] = {}
        
        return extracted_data